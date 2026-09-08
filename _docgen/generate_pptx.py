"""Week 15 final presentation — same models as the Word documents."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

import model as M

ROOT = Path(r"C:\Users\bwaly\Desktop\Group 20 Pharmacy-project\Pharmacy-POS")
D = ROOT / "Diagrams"
UI = ROOT / "UI-Prototypes"
OUT = ROOT / "Presentation" / "Pharmacy-POS-Final-Presentation.pptx"

NAVY = RGBColor(0x0D, 0x3B, 0x4C)
TEAL = RGBColor(0x1A, 0x7A, 0x6D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x23, 0x32)
MUTED = RGBColor(0x5B, 0x67, 0x70)
SAND = RGBColor(0xF4, 0xEF, 0xE6)
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
PALE = RGBColor(0xE6, 0xF3, 0xF1)

W = Inches(13.333)
H = Inches(7.5)


def _rgb_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _text(tf, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return p


def _add_runs(tf, lines, size=16, color=INK, bold=False, bullet=True):
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        if bullet:
            p.text = ""  # we'll set run
        run = p.add_run()
        run.text = (("•  " if bullet else "") + line)
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Calibri"
        run.font.bold = bold


def new_prs():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def band(slide, color=NAVY, height=Inches(0.18), top=0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, top, W, height)
    _rgb_fill(s, color)
    return s


def footer(slide, page, total=35):
    band(slide, TEAL, Inches(0.12), H - Inches(0.12))
    box = slide.shapes.add_textbox(Inches(0.5), H - Inches(0.42), Inches(10), Inches(0.28))
    _text(box.text_frame, "Group 20  ·  CSC 4630  ·  Pharmacy POS  ·  Unified Process  ·  Documentation phase",
          11, color=MUTED)
    n = slide.shapes.add_textbox(Inches(11.6), H - Inches(0.42), Inches(1.3), Inches(0.28))
    _text(n.text_frame, f"{page} / {total}", 11, color=MUTED, align=PP_ALIGN.RIGHT)


def title_bar(slide, kicker, title):
    band(slide, NAVY, Inches(1.15), 0)
    k = slide.shapes.add_textbox(Inches(0.55), Inches(0.12), Inches(12), Inches(0.32))
    _text(k.text_frame, kicker.upper(), 12, True, TEAL)
    t = slide.shapes.add_textbox(Inches(0.55), Inches(0.42), Inches(12.2), Inches(0.6))
    _text(t.text_frame, title, 28, True, WHITE)
    footer(slide, slide_no(slide))


def slide_no(slide):
    # filled later; temporary
    return getattr(slide, "_pg", 0)


def content_slide(prs, kicker, title, page):
    s = blank(prs)
    s._pg = page
    title_bar(s, kicker, title)
    return s


def add_picture_fit(slide, path, left, top, width, height):
    path = Path(path)
    if path.exists():
        slide.shapes.add_picture(str(path), left, top, width=width)


def card(slide, left, top, width, height, fill=PALE):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _rgb_fill(sh, fill)
    sh.line.color.rgb = TEAL
    sh.adjustments[0] = 0.08
    return sh


def build():
    prs = new_prs()
    page = 0

    # 1 Title
    page = 1
    s = blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _rgb_fill(bg, NAVY)
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), H)
    _rgb_fill(accent, TEAL)
    box = s.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(11.5), Inches(0.4))
    _text(box.text_frame, "CSC 4630  ·  ADVANCED SOFTWARE ENGINEERING  ·  GROUP 20", 14, True, TEAL)
    box = s.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.6))
    _text(box.text_frame, "Pharmacy Point-of-Sale System", 40, True, WHITE)
    box = s.shapes.add_textbox(Inches(0.9), Inches(3.7), Inches(11.5), Inches(0.6))
    _text(box.text_frame, "Final presentation  ·  Unified Process  ·  Analysis and design", 20, False, PALE)
    box = s.shapes.add_textbox(Inches(0.9), Inches(4.7), Inches(11.5), Inches(1.8))
    _add_runs(box.text_frame, [
        f"University: {M.UNIVERSITY}    Lecturer: {M.LECTURER}",
        "System roles: Administrator  ·  Pharmacist/Cashier (no separate Cashier)",
        "Implementation has not started — documentation and design only",
        M.DATE,
    ], size=16, color=WHITE, bullet=False)
    f = s.shapes.add_textbox(Inches(0.9), Inches(6.8), Inches(11), Inches(0.3))
    _text(f.text_frame, "1 / 35", 12, color=PALE)

    # 2 Problem
    page = 2
    s = content_slide(prs, "Weeks 1–2  ·  Inception", "Problem statement", page)
    box = s.shapes.add_textbox(Inches(0.6), Inches(1.45), Inches(12.1), Inches(2.2))
    _text(box.text_frame, M.PROBLEM_STATEMENT[:420] + "…", 16, color=INK)
    items = [
        ("Expired stock", "Paper and memory let expired packs reach a customer."),
        ("False availability", "A medicine is promised that the shelf no longer has."),
        ("Weak audit", "Nobody can show who sold or adjusted a batch."),
        ("Wrong product shape", "A supermarket till has no prescription or batch expiry."),
    ]
    for i, (h, b) in enumerate(items):
        c = card(s, Inches(0.55 + (i % 4) * 3.15), Inches(4.0), Inches(3.0), Inches(2.4))
        tb = s.shapes.add_textbox(Inches(0.7 + (i % 4) * 3.15), Inches(4.15), Inches(2.7), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True
        _text(tf, h, 16, True, NAVY)
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = b
        r.font.size = Pt(13)
        r.font.color.rgb = INK

    # 3 Business case
    page = 3
    s = content_slide(prs, "Inception", "Business case", page)
    cols = [
        ("Operational", "Accurate sale + matching batch update. Receipts that match the till."),
        ("Strategic", "Expiry control and prescription recording become ordinary work, not a later panic."),
        ("Financial", "Less wasted stock and less time reconstructing a day's takings. No invented ROI %."),
        ("Boundary", "NHIMA claims and e-commerce would be a different system. They stay out."),
    ]
    for i, (h, b) in enumerate(cols):
        card(s, Inches(0.5 + i * 3.2), Inches(1.6), Inches(3.05), Inches(4.6), SAND if i == 3 else PALE)
        tb = s.shapes.add_textbox(Inches(0.7 + i * 3.2), Inches(1.85), Inches(2.7), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True
        _text(tf, h, 18, True, TEAL)
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = "\n" + b
        r.font.size = Pt(15)
        r.font.color.rgb = INK

    # 4 Vision
    page = 4
    s = content_slide(prs, "Inception", "Vision and objectives", page)
    box = s.shapes.add_textbox(Inches(0.6), Inches(1.45), Inches(12.2), Inches(1.5))
    tf = box.text_frame
    tf.word_wrap = True
    _text(tf, M.VISION, 16, False, INK)
    for i, (oid, txt) in enumerate(M.OBJECTIVES):
        y = Inches(3.15 + (i % 4) * 0.85)
        x = Inches(0.55 if i < 4 else 6.9)
        card(s, x, y, Inches(6.15), Inches(0.75))
        tb = s.shapes.add_textbox(x + Inches(0.15), y + Inches(0.12), Inches(5.85), Inches(0.55))
        _text(tb.text_frame, f"{oid}  {txt}", 12, color=NAVY)

    # 5 Scope
    page = 5
    s = content_slide(prs, "Inception", "System scope", page)
    card(s, Inches(0.5), Inches(1.5), Inches(6.0), Inches(5.2), PALE)
    card(s, Inches(6.8), Inches(1.5), Inches(6.0), Inches(5.2), SAND)
    t1 = s.shapes.add_textbox(Inches(0.7), Inches(1.65), Inches(5.6), Inches(0.4))
    _text(t1.text_frame, "In scope", 20, True, TEAL)
    b1 = s.shapes.add_textbox(Inches(0.7), Inches(2.15), Inches(5.6), Inches(4.3))
    b1.text_frame.word_wrap = True
    _add_runs(b1.text_frame, M.SCOPE_IN[:8], size=13)
    t2 = s.shapes.add_textbox(Inches(7.0), Inches(1.65), Inches(5.6), Inches(0.4))
    _text(t2.text_frame, "Out of scope", 20, True, GOLD)
    b2 = s.shapes.add_textbox(Inches(7.0), Inches(2.15), Inches(5.6), Inches(4.3))
    b2.text_frame.word_wrap = True
    _add_runs(b2.text_frame, M.SCOPE_OUT[:8], size=13)

    # 6 Stakeholders
    page = 6
    s = content_slide(prs, "Inception", "Stakeholders and actors", page)
    note = s.shapes.add_textbox(Inches(0.55), Inches(1.4), Inches(12.2), Inches(0.7))
    tf = note.text_frame
    tf.word_wrap = True
    _text(tf, "The Pharmacist performs the cashier function; there is no separate Cashier role.", 18, True, TEAL)
    card(s, Inches(0.5), Inches(2.25), Inches(6.0), Inches(4.4))
    card(s, Inches(6.8), Inches(2.25), Inches(6.0), Inches(4.4), SAND)
    a = s.shapes.add_textbox(Inches(0.7), Inches(2.4), Inches(5.6), Inches(4.1))
    a.text_frame.word_wrap = True
    _text(a.text_frame, "System users (accounts)", 16, True, NAVY)
    _add_runs(a.text_frame, [
        "Administrator — catalogue, stock, users, suppliers, reports, settings",
        "Pharmacist/Cashier — search, sell, prescription capture, pay, receipt, void",
    ], size=15)
    # fix: _add_runs clears. rewrite properly
    a.text_frame.clear()
    _text(a.text_frame, "System users (accounts)", 16, True, NAVY)
    for line in [
        "Administrator — catalogue, stock, users, suppliers, reports, settings",
        "Pharmacist/Cashier — search, sell, prescription, payment, receipt, void",
    ]:
        p = a.text_frame.add_paragraph()
        r = p.add_run()
        r.text = "•  " + line
        r.font.size = Pt(15)
        r.font.color.rgb = INK
    b = s.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.6), Inches(4.1))
    b.text_frame.word_wrap = True
    _text(b.text_frame, "External (no login)", 16, True, NAVY)
    for line in [
        "Customer — buys; talks to the Pharmacist, not the software",
        "Supplier — source of batches; recorded on receipts of stock",
        "No Manager, Cashier, or Inventory Clerk account",
    ]:
        p = b.text_frame.add_paragraph()
        r = p.add_run()
        r.text = "•  " + line
        r.font.size = Pt(15)
        r.font.color.rgb = INK

    # 7 UP
    page = 7
    s = content_slide(prs, "Methodology", "Unified Process — iterative, not waterfall", page)
    add_picture_fit(s, D / "UP-Planning" / "UP-Phases.png", Inches(0.4), Inches(1.35), Inches(12.5), Inches(5.6))

    # 8 Plan
    page = 8
    s = content_slide(prs, "Methodology", "15-week iteration plan", page)
    add_picture_fit(s, D / "UP-Planning" / "Iteration-Plan-15-Weeks.png", Inches(0.35), Inches(1.3), Inches(12.6), Inches(5.7))

    # 9 10% UC
    page = 9
    s = content_slide(prs, "Inception  ·  10%", "Initial use cases — Process Sale", page)
    add_picture_fit(s, D / "Use-Cases" / "UC-Inception-10-Percent.png", Inches(0.35), Inches(1.3), Inches(12.6), Inches(5.7))

    # 10 30%
    page = 10
    s = content_slide(prs, "Elaboration 1  ·  30%", "Architecturally significant use cases", page)
    add_picture_fit(s, D / "Use-Cases" / "UC-Elaboration1-30-Percent.png", Inches(0.35), Inches(1.3), Inches(12.6), Inches(5.7))

    # 11 Detailed UC
    page = 11
    s = content_slide(prs, "Elaboration 1", "Detailed use case — UC-02 Process Sale", page)
    left = s.shapes.add_textbox(Inches(0.55), Inches(1.4), Inches(6.1), Inches(5.5))
    tf = left.text_frame
    tf.word_wrap = True
    _text(tf, "Main success scenario (short)", 16, True, TEAL)
    for line in [
        "makeNewSale — empty sale in progress",
        "enterItem(medicineId, qty) — repeat",
        "Refuse expired or short batches",
        "attachPrescription when required",
        "endSale — show total",
        "makePayment(amount, method)",
        "Stock down + receipt out",
    ]:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + line
        r.font.size = Pt(15)
        r.font.color.rgb = INK
    right = s.shapes.add_textbox(Inches(6.9), Inches(1.4), Inches(5.9), Inches(5.5))
    tf = right.text_frame
    tf.word_wrap = True
    _text(tf, "Why this use case leads design", 16, True, TEAL)
    for line in [
        "Touches money, stock, expiry, Rx, audit",
        "Customer does not generate system events",
        "Scan is only a technology of enterItem",
        "Fully dressed text is in the Elaboration 1 report",
        "Extensions cover short cash and failed persist",
    ]:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + line
        r.font.size = Pt(15)
        r.font.color.rgb = INK

    # 12 SSD
    page = 12
    s = content_slide(prs, "Elaboration 1", "System sequence diagram — Process Sale", page)
    add_picture_fit(s, D / "SSD" / "SSD-UC02-Process-Sale.png", Inches(0.35), Inches(1.28), Inches(12.6), Inches(5.75))

    # 13 Domain
    page = 13
    s = content_slide(prs, "Elaboration 1", "Domain model (conceptual — no methods)", page)
    add_picture_fit(s, D / "Domain-Model" / "Domain-Model.png", Inches(0.2), Inches(1.25), Inches(12.9), Inches(5.8))

    # 14 Architecture
    page = 14
    s = content_slide(prs, "Elaboration 1", "Architectural proof-of-concept", page)
    add_picture_fit(s, D / "Architecture" / "Layered-Architecture-PoC.png", Inches(0.3), Inches(1.28), Inches(12.7), Inches(5.75))

    # 15 Elab 2
    page = 15
    s = content_slide(prs, "Elaboration 2", "From requirements to design", page)
    steps = [
        ("Use cases 70%", "UC-04 to UC-10 dressed"),
        ("SSDs", "Inventory and void events"),
        ("Responsibilities", "GRASP on each sale operation"),
        ("DCD", "Methods on software classes"),
        ("UI", "Both roles, no third role"),
        ("Schema", "3NF + sale price copy"),
    ]
    for i, (h, b) in enumerate(steps):
        r, c = divmod(i, 3)
        card(s, Inches(0.5 + c * 4.2), Inches(1.7 + r * 2.4), Inches(4.0), Inches(2.15))
        tb = s.shapes.add_textbox(Inches(0.7 + c * 4.2), Inches(1.9 + r * 2.4), Inches(3.6), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        _text(tf, h, 18, True, NAVY)
        p = tf.add_paragraph()
        run = p.add_run()
        run.text = b
        run.font.size = Pt(15)
        run.font.color.rgb = INK

    # 16 Complete UC
    page = 16
    s = content_slide(prs, "Elaboration 2", "Complete use-case model", page)
    add_picture_fit(s, D / "Use-Cases" / "UC-Complete.png", Inches(0.25), Inches(1.25), Inches(12.8), Inches(5.8))

    # 17 DCD
    page = 17
    s = content_slide(prs, "Elaboration 2", "Design class diagram — Process Sale", page)
    add_picture_fit(s, D / "Design-Class" / "DCD-Process-Sale.png", Inches(0.2), Inches(1.25), Inches(12.9), Inches(5.8))

    # 18 GRASP
    page = 18
    s = content_slide(prs, "Elaboration 2", "GRASP — who does what, and why", page)
    add_picture_fit(s, D / "Design-Class" / "GRASP-Assignments.png", Inches(0.3), Inches(1.28), Inches(12.7), Inches(5.75))

    # 19 UI
    page = 19
    s = content_slide(prs, "Elaboration 2", "UI prototype — Pharmacist counter", page)
    add_picture_fit(s, UI / "UI-02-POS-Sale.png", Inches(0.35), Inches(1.3), Inches(12.6), Inches(5.7))

    # 19b we only have 35 slides - admin UI will appear in final system slide

    # 20 DB design
    page = 20
    s = content_slide(prs, "Elaboration 2", "Database design groups", page)
    add_picture_fit(s, D / "Database" / "Schema-Groups.png", Inches(0.35), Inches(1.3), Inches(12.6), Inches(5.7))

    # 21 ERD
    page = 21
    s = content_slide(prs, "Elaboration 2", "Entity-relationship diagram", page)
    add_picture_fit(s, D / "Database" / "ERD-Logical.png", Inches(0.15), Inches(1.22), Inches(13.0), Inches(5.85))

    # 22 Schema
    page = 22
    s = content_slide(prs, "Elaboration 2", "Major tables and relationships", page)
    body = s.shapes.add_textbox(Inches(0.6), Inches(1.45), Inches(12.2), Inches(5.5))
    tf = body.text_frame
    tf.word_wrap = True
    _text(tf, "Persistent form of the domain model — design only, no SQL scripts", 16, True, TEAL)
    lines = [
        "users.role ∈ {ADMINISTRATOR, PHARMACIST}  ·  password_hash never plaintext",
        "medicines described; stock_batches hold expiry and quantity_on_hand",
        "sales 1—* sale_items  ·  sale_items copies unit_price (history stays still)",
        "sales 1—0..1 payments  ·  method = CASH | CARD | MOBILE_MONEY",
        "customers / prescriptions support attachPrescription; no customer login",
        "stock_receipts + stock_adjustments are Administrator movements",
        "audit_logs record login, sale, void, stock, and user-admin events",
        "3NF target; M:N resolved by sale_items and prescription_items",
    ]
    for line in lines:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + line
        r.font.size = Pt(16)
        r.font.color.rgb = INK

    # 23 Construction
    page = 23
    s = content_slide(prs, "Weeks 9–12", "Construction — Iterations 3 and 4", page)
    card(s, Inches(0.5), Inches(1.55), Inches(6.0), Inches(5.1))
    card(s, Inches(6.8), Inches(1.55), Inches(6.0), Inches(5.1), SAND)
    t = s.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(5.6), Inches(4.8))
    tf = t.text_frame
    tf.word_wrap = True
    _text(tf, "Iteration 3 — spine", 18, True, NAVY)
    for line in ["F-01 Auth", "F-08 Medicines", "F-02…F-07 Process Sale", "UT-01–UT-09  ·  IT-01–IT-03"]:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + line
        r.font.size = Pt(16)
    t = s.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.6), Inches(4.8))
    tf = t.text_frame
    tf.word_wrap = True
    _text(tf, "Iteration 4 — widen", 18, True, NAVY)
    for line in ["F-09–F-10 Inventory", "F-15 Void", "F-11–F-16 Admin + reports", "Remaining tests"]:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + line
        r.font.size = Pt(16)
    warn = s.shapes.add_textbox(Inches(0.6), Inches(6.35), Inches(12), Inches(0.4))
    _text(warn.text_frame, "Feature status is Planned. Results: TO BE COMPLETED DURING CONSTRUCTION.", 14, True, GOLD)

    # 24 Testing
    page = 24
    s = content_slide(prs, "Construction", "Testing strategy", page)
    card(s, Inches(0.5), Inches(1.55), Inches(6.0), Inches(5.1))
    card(s, Inches(6.8), Inches(1.55), Inches(6.0), Inches(5.1))
    t = s.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(5.6), Inches(4.8))
    tf = t.text_frame
    tf.word_wrap = True
    _text(tf, "Unit (examples)", 18, True, TEAL)
    for line in ["UT-04 expired batch not sellable", "UT-08 failed persist leaves stock", "UT-10 void restores quantity", "UT-02 wrong password → no session"]:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + line
        r.font.size = Pt(15)
    t = s.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.6), Inches(4.8))
    tf = t.text_frame
    tf.word_wrap = True
    _text(tf, "Integration (examples)", 18, True, TEAL)
    for line in ["IT-01 full OTC path in the database", "IT-02 expiry refusal leaves no line", "IT-05 void + audit row", "IT-06 Pharmacist cannot create medicines"]:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + line
        r.font.size = Pt(15)

    # 25 Git CI
    page = 25
    s = content_slide(prs, "Construction", "Git and continuous integration", page)
    add_picture_fit(s, D / "UP-Planning" / "Git-CI-Workflow.png", Inches(0.4), Inches(1.3), Inches(12.5), Inches(5.7))

    # 26 Defects
    page = 26
    s = content_slide(prs, "Construction", "Defect tracking", page)
    body = s.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.4))
    tf = body.text_frame
    tf.word_wrap = True
    _text(tf, "A defect is a failed test or a broken use-case guarantee — not a new feature request.", 16, True, NAVY)
    for line in [
        "Log fields: DEF-id, test id, UC/FR, severity, status, commit hash",
        "Template row DEF-000 exists so the log format is ready",
        "No invented defect counts or “12 bugs closed” stories",
        "New ideas (NHIMA, web shop) go to the backlog, not the defect log",
        "Status: TO BE COMPLETED DURING CONSTRUCTION",
    ]:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + line
        r.font.size = Pt(17)
        r.font.color.rgb = INK

    # 27 Transition
    page = 27
    s = content_slide(prs, "Weeks 13–14", "Transition — beta, performance, security", page)
    items = [
        ("Beta users", "One Administrator account, one Pharmacist/Cashier account. Same human may switch — never a merged role."),
        ("Feedback", "Structured prompts. Analysis splits usability / defect / out-of-scope."),
        ("Performance", "Time enterItem on the real workstation against NFR-02. No fake timings today."),
        ("Security", "Plaintext password must fail; Pharmacist must not open user-create; voids need reasons."),
    ]
    for i, (h, b) in enumerate(items):
        card(s, Inches(0.5 + (i % 2) * 6.4), Inches(1.55 + (i // 2) * 2.5), Inches(6.15), Inches(2.3))
        tb = s.shapes.add_textbox(Inches(0.7 + (i % 2) * 6.4), Inches(1.7 + (i // 2) * 2.5), Inches(5.8), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True
        _text(tf, h, 18, True, TEAL)
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = b
        r.font.size = Pt(14)

    # 28 Final system
    page = 28
    s = content_slide(prs, "Designed system", "What the constructed system will be", page)
    add_picture_fit(s, UI / "UI-06-Admin-Dashboard.png", Inches(0.3), Inches(1.35), Inches(6.3), Inches(3.6))
    add_picture_fit(s, UI / "UI-03-Payment.png", Inches(6.7), Inches(1.35), Inches(6.2), Inches(3.6))
    cap = s.shapes.add_textbox(Inches(0.5), Inches(5.1), Inches(12.3), Inches(1.5))
    tf = cap.text_frame
    tf.word_wrap = True
    _text(tf, "Administrator dashboard and Pharmacist payment — same two roles on every screen. "
          "The running product will replace these prototypes after construction.", 16, False, INK)

    # 29 Final UML
    page = 29
    s = content_slide(prs, "Week 15", "Final UML — the same approved models", page)
    add_picture_fit(s, D / "Architecture" / "Architecture-Sale-Path.png", Inches(0.3), Inches(1.3), Inches(12.7), Inches(5.7))

    # 30 Reflection
    page = 30
    s = content_slide(prs, "Week 15", "UP reflection", page)
    refs_txt = [
        ("Inception", "Treating 10% as Process Sale stopped insurance/refill scope creep from an earlier briefing slide."),
        ("Elaboration", "SSDs and StockBatch made the pharmacy difference from NextGen POS real. Register was omitted on purpose."),
        ("Construction", "Not lived yet. Planned lesson: vertical slices, not layer-by-layer waterfall."),
        ("Transition", "Not lived yet. R-05 (training) will be confirmed only when a non-author Pharmacist uses the till."),
    ]
    for i, (h, b) in enumerate(refs_txt):
        card(s, Inches(0.5), Inches(1.45 + i * 1.3), Inches(12.3), Inches(1.18))
        tb = s.shapes.add_textbox(Inches(0.7), Inches(1.55 + i * 1.3), Inches(12.0), Inches(1.0))
        tf = tb.text_frame
        tf.word_wrap = True
        _text(tf, h, 16, True, TEAL)
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = b
        r.font.size = Pt(14)

    # 31 Lessons
    page = 31
    s = content_slide(prs, "Week 15", "Lessons learned (design phase)", page)
    lessons = [
        "A use case is a goal. “Click Search” is not a use case.",
        "Two roles change the diagram, the UI, and the users table — not just a sentence on a slide.",
        "Pharmacy POS ≠ NextGen POS. Expiry and prescriptions in; Register out.",
        "Assumptions that are not written become fake requirements.",
        "“Mitigated in design” is not “gone”.",
        "Stable IDs (FR-05, UC-02, UT-04) are how UP documents stay one project.",
    ]
    body = s.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(12), Inches(5.4))
    tf = body.text_frame
    tf.word_wrap = True
    _text(tf, "", 12)
    tf.clear()
    for i, line in enumerate(lessons):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + line
        r.font.size = Pt(18)
        r.font.color.rgb = INK
        p.space_after = Pt(10)

    # 32 Achievements
    page = 32
    s = content_slide(prs, "Week 15", "Key achievements (this pack)", page)
    ach = [
        "Marking-guide compliance matrix covering all 100 marks",
        "Ten goal-level use cases with a justified 10 / 30 / 70 split",
        "Larman-style SSDs, domain model, DCDs, and GRASP",
        "Logical schema + ERD consistent with Process Sale",
        "UI prototypes for both roles only",
        "Honest placeholders where construction has not occurred",
    ]
    for i, line in enumerate(ach):
        card(s, Inches(0.55 + (i % 2) * 6.35), Inches(1.55 + (i // 2) * 1.7), Inches(6.15), Inches(1.5))
        tb = s.shapes.add_textbox(Inches(0.75 + (i % 2) * 6.35), Inches(1.8 + (i // 2) * 1.7), Inches(5.8), Inches(1.15))
        tb.text_frame.word_wrap = True
        _text(tb.text_frame, line, 16, False, NAVY)

    # 33 Demo plan
    page = 33
    s = content_slide(prs, "Week 15", "Live demonstration plan", page)
    demo = [
        ("SC-01", "OTC cash sale → receipt → stock down"),
        ("SC-02", "Prescription-only line blocked until details exist"),
        ("SC-03", "Expired batch refused"),
        ("SC-06", "Void restores stock and keeps the sale"),
        ("SC-07", "Pharmacist cannot open user administration"),
        ("SC-08", "Administrator report shows the completed sale once"),
    ]
    for i, (h, b) in enumerate(demo):
        card(s, Inches(0.5 + (i % 3) * 4.2), Inches(1.6 + (i // 3) * 2.4), Inches(4.0), Inches(2.15))
        tb = s.shapes.add_textbox(Inches(0.7 + (i % 3) * 4.2), Inches(1.8 + (i // 3) * 2.4), Inches(3.6), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        _text(tf, h, 20, True, TEAL)
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = b
        r.font.size = Pt(15)

    # 34 Conclusion
    page = 34
    s = content_slide(prs, "Close", "Conclusion", page)
    body = s.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    tf = body.text_frame
    tf.word_wrap = True
    _text(tf, "Group 20 has a single connected design.", 22, True, NAVY)
    for line in [
        "Process Sale is the thread through every artefact.",
        "Exactly two roles — the Pharmacist is the cashier.",
        "UP phases and the lecturer’s 15-week marks are mapped, not decorated.",
        "Code starts only on the instruction START IMPLEMENTATION.",
    ]:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = "\n•  " + line
        r.font.size = Pt(20)
        r.font.color.rgb = INK

    # 35 Questions
    page = 35
    s = blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _rgb_fill(bg, NAVY)
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), H)
    _rgb_fill(accent, TEAL)
    box = s.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.2))
    _text(box.text_frame, "Questions", 48, True, WHITE)
    box = s.shapes.add_textbox(Inches(0.9), Inches(3.8), Inches(11.5), Inches(1.4))
    _text(box.text_frame, "Group 20  ·  CSC 4630  ·  Pharmacy POS\nUnified Process  ·  Documentation first", 20, False, PALE)
    box = s.shapes.add_textbox(Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.4))
    _text(box.text_frame, "35 / 35", 12, color=PALE)

    # Fix footers that used slide_no 0 — title_bar already wrote 0 for content slides
    # Re-open is hard; we encoded page in title_bar via s._pg. Recreate footers by
    # setting page numbers at creation — we passed page into content_slide. title_bar
    # calls slide_no which reads _pg. Good if _pg is set before title_bar... 
    # content_slide sets _pg then title_bar. Good.

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT, "slides", len(prs.slides))
    return OUT


if __name__ == "__main__":
    build()
