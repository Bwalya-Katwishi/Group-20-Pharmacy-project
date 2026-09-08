"""Shared Word formatting for Pharmacy POS documentation."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Inches, Pt, RGBColor, Emu
from docx.enum.style import WD_STYLE_TYPE

NAVY = RGBColor(0x0D, 0x3B, 0x4C)
TEAL = RGBColor(0x1A, 0x7A, 0x6D)
INK = RGBColor(0x1A, 0x23, 0x32)
MUTED = RGBColor(0x5B, 0x67, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
HEADER_BG = "0D3B4C"
ALT_BG = "E6F3F1"
PLACEHOLDER_BG = "F4EFE6"

ROOT = Path(r"C:\Users\bwaly\Desktop\Group 20 Pharmacy-project\Pharmacy-POS")


def _set_run_font(run, name="Calibri", size=11, bold=False, color=INK, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def _shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), kwargs.get("val", "single"))
        el.set(qn("w:sz"), kwargs.get("sz", "4"))
        el.set(qn("w:color"), kwargs.get("color", "C5D0D4"))
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def _add_toc_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-3" \h \z \u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Right-click and choose Update Field in Microsoft Word to refresh the table of contents."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(text)
    r.append(fld_end)


def new_document(header_left, header_right):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.clear()
    r = hp.add_run(f"{header_left}    |    {header_right}")
    _set_run_font(r, size=9, color=TEAL)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.clear()
    r1 = fp.add_run("Group 20  ·  CSC 4630  ·  Unified Process  ·  Page ")
    _set_run_font(r1, size=9, color=MUTED)
    _add_page_number(fp)
    extra = fp.add_run("  ·  Documentation only — no implementation yet")
    _set_run_font(extra, size=9, color=MUTED)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].font.color.rgb = INK
    styles["Normal"].paragraph_format.space_after = Pt(8)
    styles["Normal"].paragraph_format.line_spacing = 1.15

    for i, (size, color) in enumerate([(22, NAVY), (16, TEAL), (13, NAVY)], start=1):
        st = styles[f"Heading {i}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(16 if i == 1 else 12)
        st.paragraph_format.space_after = Pt(8)

    return doc


def cover(doc, kicker, title, subtitle, extra_lines):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(kicker.upper())
    _set_run_font(r, size=12, bold=True, color=TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    _set_run_font(r, size=28, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    _set_run_font(r, size=14, color=TEAL)

    doc.add_paragraph()
    for line in extra_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        _set_run_font(r, size=12, color=INK)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run(
        "\nAnalysis and design artefact. No application source code is included. "
        "Implementation starts only after the instruction START IMPLEMENTATION."
    )
    _set_run_font(r, size=10, italic=True, color=MUTED)
    doc.add_page_break()


def toc(doc):
    h = doc.add_heading("Table of Contents", level=1)
    p = doc.add_paragraph()
    _add_toc_field(p)
    hint = doc.add_paragraph()
    r = hint.add_run("In Word: References → Update Table, or right-click the field above.")
    _set_run_font(r, size=9, italic=True, color=MUTED)
    doc.add_page_break()


def h1(doc, text):
    return doc.add_heading(text, level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def h3(doc, text):
    return doc.add_heading(text, level=3)


def para(doc, text, *, bold=False, italic=False, color=INK, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    p.paragraph_format.space_after = Pt(8)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    for run in p.runs:
        _set_run_font(run, size=11)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(text, style="List Number")
    for run in p.runs:
        _set_run_font(run, size=11)
    return p


def placeholder(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run_font(r, size=11, bold=True, italic=True, color=RGBColor(0x8A, 0x5A, 0x00))
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run_font(r, size=10, italic=True, color=TEAL)
    p.paragraph_format.space_after = Pt(4)
    return p


def explain(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run_font(r, size=10, italic=False, color=MUTED)
    p.paragraph_format.space_after = Pt(12)
    return p


def add_figure(doc, path, caption_text, explanation, width=15.8):
    path = Path(path)
    if not path.exists():
        para(doc, f"[Missing figure: {path.name}]", italic=True, color=MUTED)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width))
    caption(doc, caption_text)
    explain(doc, explanation)


def table(doc, headers, rows, col_widths=None, header_hex=HEADER_BG):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        _set_run_font(r, size=9, bold=True, color=WHITE)
        _shade(cell, header_hex)
        _set_cell_border(cell, color="0D3B4C")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            _set_run_font(r, size=9, color=INK)
            if ri % 2 == 0:
                _shade(cell, ALT_BG)
            _set_cell_border(cell)
    if col_widths:
        for row in t.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def refs(doc):
    h1(doc, "References")
    items = [
        "Larman, C. Applying UML and Patterns: An Introduction to Object-Oriented Analysis and Design and Iterative Development. 3rd edition. Prentice Hall. Principal methodology for UP, use cases, domain models, SSDs, design class diagrams, GRASP, and iterative development. Page numbers are not cited because the full print edition was not available as a continuous source file; chapter titles from the course extracts (System Sequence Diagrams; Domain Model associations; Design Class Diagrams) are used instead.",
        "CSC 4630 Advanced Software Engineering Course Project Marking Guide. Unified Process, 15 weeks, 100 marks. Determines milestones, deliverables, and mark weights.",
        "Larman course extracts used by the group: Chapter 9 System Sequence Diagrams; Chapter 11 Domain Model associations; Chapter 16 Design Class Diagrams.",
        "Group 20 inception briefing slides, Pharmacy POS System — Advanced Software Engineering (problem context in Zambia). Role names in those slides that implied a separate Cashier or Manager were corrected in this documentation to match the two-role rule.",
    ]
    for i, item in enumerate(items, 1):
        para(doc, f"[{i}]  {item}")


def dressed_use_case(doc, uc, detail):
    h3(doc, f"{uc['id']}  {uc['name']}")
    table(doc,
          ["Field", "Description"],
          [
              ["ID", uc["id"]],
              ["Name", uc["name"]],
              ["Scope", detail["scope"]],
              ["Level", detail["level"]],
              ["Primary actor", detail["primary"]],
              ["Priority", uc["priority"]],
              ["Iteration", uc["iteration"]],
              ["Requirements", ", ".join(uc["frs"])],
              ["Preconditions", detail["preconditions"]],
              ["Success guarantee", detail["success"]],
              ["Special requirements", detail["special"]],
              ["Technology / data variations", detail["tech"]],
              ["Frequency", detail["frequency"]],
          ])
    para(doc, "Stakeholders and interests", bold=True)
    table(doc, ["Stakeholder", "Interest"], detail["stakeholders"])
    para(doc, "Main success scenario", bold=True)
    for i, step in enumerate(detail["main"], 1):
        para(doc, f"{i}.  {step}")
    para(doc, "Extensions", bold=True)
    for ext in detail["extensions"]:
        bullet(doc, ext)
