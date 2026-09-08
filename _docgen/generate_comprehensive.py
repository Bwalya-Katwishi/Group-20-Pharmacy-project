"""Merge all UP phase Word documents into one comprehensive submission report."""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

import model as M
from word_theme import ROOT, cover, new_document, h1, para, refs

OUT_DIR = ROOT / "Documentation" / "09-Comprehensive-Report"
OUT = OUT_DIR / "Pharmacy-POS-Comprehensive-Project-Report.docx"

META = [
    f"{M.COURSE}",
    f"{M.GROUP}",
    f"Lecturer: {M.LECTURER}",
    f"University: {M.UNIVERSITY}",
    f"{M.DATE}",
    M.METHODOLOGY,
]

PARTS = [
    ("00  Project Control — Marking Compliance Matrix",
     ROOT / "Documentation" / "00-Project-Control" / "Marking-Guide-Compliance-Matrix.docx"),
    ("01  Inception (Weeks 1–2)",
     ROOT / "Documentation" / "01-Inception" / "Pharmacy-POS-Inception-Report.docx"),
    ("02  Elaboration Iteration 1 (Weeks 3–5)",
     ROOT / "Documentation" / "02-Elaboration-Iteration-1" / "Pharmacy-POS-Elaboration-Iteration-1-Report.docx"),
    ("03  Elaboration Iteration 2 (Weeks 6–8)",
     ROOT / "Documentation" / "03-Elaboration-Iteration-2" / "Pharmacy-POS-Elaboration-Iteration-2-Report.docx"),
    ("04  Complete System Design",
     ROOT / "Documentation" / "04-Complete-System-Design" / "Pharmacy-POS-Complete-System-Design.docx"),
    ("05  Construction & Testing (Weeks 9–12)",
     ROOT / "Documentation" / "05-Construction-Testing" / "Pharmacy-POS-Construction-Testing-Report.docx"),
    ("06  Transition & Beta (Weeks 13–14)",
     ROOT / "Documentation" / "06-Transition-Beta" / "Pharmacy-POS-Transition-Beta-Report.docx"),
    ("07  User Manual",
     ROOT / "Documentation" / "07-User-Manual" / "Pharmacy-POS-User-Manual.docx"),
    ("08  Database Design",
     ROOT / "Database-Design" / "Pharmacy-POS-Database-Design.docx"),
    ("09  Final Project Report (Week 15)",
     ROOT / "Documentation" / "08-Final-Report" / "Pharmacy-POS-Final-Project-Report.docx"),
    ("10  Diagram Inventory",
     ROOT / "Documentation" / "00-Project-Control" / "Diagram-Inventory.docx"),
]


def _append_body(master: Document, slave: Document, page_break: bool = True) -> None:
    if page_break:
        p = master.add_paragraph()
        run = p.add_run()
        run.add_break(WD_BREAK.PAGE)
    for element in slave.element.body:
        master.element.body.append(deepcopy(element))


def build() -> Path:
    missing = [str(p) for _, p in PARTS if not p.exists()]
    if missing:
        raise FileNotFoundError(
            "Run generate_docs_part1.py and generate_docs_part2.py first. Missing:\n  "
            + "\n  ".join(missing)
        )

    master = new_document(
        "Pharmacy POS  ·  Comprehensive Report",
        "All UP deliverables — single submission volume",
    )
    cover(
        master,
        "CSC 4630  ·  Group 20",
        "Comprehensive Project Report",
        "Pharmacy Point-of-Sale — Inception through Transition (all phase documents)",
        META,
    )

    h1(master, "Document Index")
    para(master, (
        "This volume merges every separate deliverable required by the CSC 4630 marking guide "
        "into one file for submission and archival. Each part retains its original cover page, "
        "table of contents, and section numbering from the phase report."
    ))
    for label, path in PARTS:
        para(master, f"•  {label}  —  {path.name}")

    p = master.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)

    for i, (label, path) in enumerate(PARTS):
        slave = Document(str(path))
        _append_body(master, slave, page_break=(i > 0))

    refs(master)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    master.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
